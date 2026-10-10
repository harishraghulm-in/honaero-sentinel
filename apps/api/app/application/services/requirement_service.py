import re
import csv
import json
import io
import base64
import hashlib
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass

from apps.api.app.domain.enums import RequirementType


AMBIGUOUS_KEYWORDS = [
    "appropriate",
    "as needed",
    "as required",
    "adequate",
    "etc",
    "user friendly",
    "rapidly",
    "normal",
    "if possible",
    "sufficient",
    "optimize",
    "robust",
    "flexible",
    "timely",
]


@dataclass
class ExtractedRequirementDTO:
    identifier: str
    title: str
    description: str
    req_type: RequirementType = RequirementType.HLR
    section: Optional[str] = None
    page_or_line: Optional[str] = None
    acceptance_criteria: Optional[str] = None
    verification_method: str = "TEST"
    ambiguity_status: str = "CLEAR"
    ambiguity_notes: Optional[str] = None


@dataclass
class CandidateCaseDTO:
    name: str
    case_category: str  # NORMAL, BOUNDARY, INVALID_INPUT, CONDITION_BASED
    rationale: str
    target_function_name: Optional[str]
    preconditions: Dict[str, Any]
    input_vectors: List[Dict[str, Any]]
    expected_outputs: Dict[str, Any]
    is_expected_result_uncertain: bool
    uncertainty_reason: Optional[str]
    provenance: str = "deterministic_boundary_generator"


class RequirementService:
    """Deterministic extractor and candidate test case generator for DO-178C requirements."""

    def extract_from_document(
        self,
        content: str,
        filename: str,
        file_type: Optional[str] = None,
    ) -> List[ExtractedRequirementDTO]:
        if not content or not content.strip():
            raise ValueError("Empty requirement document: Content cannot be empty.")

        ext = filename.lower().split(".")[-1] if "." in filename else ""
        if ext in ["doc", "rtf", "odt"]:
            raise ValueError(
                f"Unsupported legacy format .{ext}. Supported formats are .txt, .md, .json, .csv, .pdf, .docx."
            )

        if ext == "pdf" or (file_type and "pdf" in file_type):
            return self._extract_from_pdf(content)
        elif ext == "docx" or (file_type and "docx" in file_type):
            return self._extract_from_docx(content)
        elif ext == "json" or (file_type and "json" in file_type):
            return self._extract_from_json(content)
        elif ext == "csv" or (file_type and "csv" in file_type):
            return self._extract_from_csv(content)
        elif ext in ["txt", "md", "markdown", ""] or (file_type and any(t in file_type for t in ["text", "markdown"])):
            return self._extract_from_text_or_markdown(content)
        else:
            raise ValueError(
                f"Unsupported document format .{ext}. Supported formats are .txt, .md, .json, .csv, .pdf, .docx."
            )

    def _get_bytes_from_content(self, content: str) -> bytes:
        try:
            decoded = base64.b64decode(content, validate=True)
            if len(decoded) > 0:
                return decoded
        except Exception:
            pass
        return content.encode("latin-1")

    def _extract_from_pdf(self, content: str) -> List[ExtractedRequirementDTO]:
        try:
            import pypdf
        except ImportError:
            raise ValueError("PDF parser unavailable in runtime environment.")

        raw_bytes = self._get_bytes_from_content(content)
        try:
            reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
        except Exception as e:
            raise ValueError(f"Malformed or corrupt PDF document: {str(e)}")

        if len(reader.pages) == 0:
            raise ValueError("Empty PDF document: No pages found.")

        results: List[ExtractedRequirementDTO] = []
        for page_idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                continue
            page_reqs = self._extract_from_text_or_markdown(
                text,
                default_section=f"PDF Page {page_idx}",
                page_prefix=f"Page {page_idx}",
            )
            results.extend(page_reqs)

        if not results:
            raise ValueError("No extractable requirements found in the uploaded PDF document.")
        return results

    def _extract_from_docx(self, content: str) -> List[ExtractedRequirementDTO]:
        try:
            import docx
        except ImportError:
            raise ValueError("DOCX parser unavailable in runtime environment.")

        raw_bytes = self._get_bytes_from_content(content)
        try:
            doc = docx.Document(io.BytesIO(raw_bytes))
        except Exception as e:
            raise ValueError(f"Malformed or corrupt DOCX document: {str(e)}")

        full_text_lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                if p.style and hasattr(p.style, "name") and "Heading" in p.style.name:
                    full_text_lines.append(f"# {p.text.strip()}")
                else:
                    full_text_lines.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text_lines.append(row_text)

        full_text = "\n\n".join(full_text_lines)
        if not full_text.strip():
            raise ValueError("Empty DOCX document: No readable text found.")

        results = self._extract_from_text_or_markdown(
            full_text,
            default_section="Document Body",
            page_prefix="Docx Section",
        )
        if not results:
            raise ValueError("No extractable requirements found in the uploaded DOCX document.")
        return results

    def _extract_from_json(self, content: str) -> List[ExtractedRequirementDTO]:
        try:
            data = json.loads(content)
        except Exception as e:
            raise ValueError(f"Malformed JSON requirement document: {str(e)}")

        items = data if isinstance(data, list) else data.get("requirements", [])
        if not isinstance(items, list):
            raise ValueError("JSON requirement document must contain a list of requirements.")

        results: List[ExtractedRequirementDTO] = []
        for idx, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                continue
            identifier = item.get("identifier") or item.get("id") or f"REQ-{idx:03d}"
            title = item.get("title") or f"Requirement {identifier}"
            description = item.get("description") or item.get("text") or ""
            req_type_str = item.get("req_type", "HLR").upper()
            req_type = RequirementType.LLR if req_type_str == "LLR" else RequirementType.HLR
            section = item.get("section")
            page_or_line = str(item.get("page_or_line", f"Item {idx}"))
            acceptance = item.get("acceptance_criteria")
            ver_method = item.get("verification_method", "TEST")

            ambiguity_status, ambiguity_notes = self._assess_ambiguity(description, acceptance)

            results.append(ExtractedRequirementDTO(
                identifier=identifier,
                title=title,
                description=description,
                req_type=req_type,
                section=section,
                page_or_line=page_or_line,
                acceptance_criteria=acceptance,
                verification_method=ver_method,
                ambiguity_status=ambiguity_status,
                ambiguity_notes=ambiguity_notes,
            ))
        return results

    def _extract_from_csv(self, content: str) -> List[ExtractedRequirementDTO]:
        results: List[ExtractedRequirementDTO] = []
        reader = csv.DictReader(io.StringIO(content))
        for idx, row in enumerate(reader, start=1):
            identifier = row.get("identifier") or row.get("id") or f"REQ-{idx:03d}"
            title = row.get("title") or f"Requirement {identifier}"
            description = row.get("description") or row.get("text") or ""
            req_type_str = row.get("req_type", "HLR").upper()
            req_type = RequirementType.LLR if req_type_str == "LLR" else RequirementType.HLR
            section = row.get("section")
            acceptance = row.get("acceptance_criteria")
            ver_method = row.get("verification_method", "TEST")

            ambiguity_status, ambiguity_notes = self._assess_ambiguity(description, acceptance)

            results.append(ExtractedRequirementDTO(
                identifier=identifier,
                title=title,
                description=description,
                req_type=req_type,
                section=section,
                page_or_line=f"Row {idx + 1}",
                acceptance_criteria=acceptance,
                verification_method=ver_method,
                ambiguity_status=ambiguity_status,
                ambiguity_notes=ambiguity_notes,
            ))
        return results

    def _extract_from_text_or_markdown(
        self,
        content: str,
        default_section: str = "General",
        page_prefix: str = "Line",
    ) -> List[ExtractedRequirementDTO]:
        lines = content.splitlines()
        results: List[ExtractedRequirementDTO] = []

        current_section = default_section
        current_id: Optional[str] = None
        current_title: Optional[str] = None
        current_lines: List[str] = []
        current_line_start = 1

        req_pattern = re.compile(r"^(?:#+\s*)?(?:\[([A-Z0-9_-]+)\]|([A-Z0-9_-]+):?)\s*(.*)$")
        section_pattern = re.compile(r"^#+\s*(?:Section\s+)?([0-9.]+\s*.*)$", re.IGNORECASE)

        for line_idx, line in enumerate(lines, start=1):
            stripped = line.strip()
            if not stripped:
                continue

            sec_match = section_pattern.match(stripped)
            if sec_match and not any(p in stripped.upper() for p in ["REQ-", "HLR-", "LLR-"]):
                current_section = sec_match.group(1).strip()
                continue

            # Check for requirement header
            req_id_match = re.search(r"\b((?:HLR|REQ|LLR|SRS|SYS)(?:-[A-Za-z0-9_]+)+)\b", stripped, re.IGNORECASE)
            if req_id_match and (stripped.startswith("#") or ":" in stripped or "-" in stripped):
                # Flush previous
                if current_id and current_lines:
                    desc = "\n".join(current_lines).strip()
                    amb_status, amb_notes = self._assess_ambiguity(desc, None)
                    results.append(ExtractedRequirementDTO(
                        identifier=current_id,
                        title=current_title or current_id,
                        description=desc,
                        req_type=RequirementType.LLR if "LLR" in current_id.upper() else RequirementType.HLR,
                        section=current_section,
                        page_or_line=f"{page_prefix} {current_line_start}",
                        acceptance_criteria=self._extract_acceptance_criteria(desc),
                        verification_method="TEST",
                        ambiguity_status=amb_status,
                        ambiguity_notes=amb_notes,
                    ))

                current_id = req_id_match.group(1).upper()
                current_title = stripped.lstrip("#").strip()
                current_lines = [stripped]
                current_line_start = line_idx
            else:
                if current_id:
                    current_lines.append(stripped)

        # Flush final requirement
        if current_id and current_lines:
            desc = "\n".join(current_lines).strip()
            amb_status, amb_notes = self._assess_ambiguity(desc, None)
            results.append(ExtractedRequirementDTO(
                identifier=current_id,
                title=current_title or current_id,
                description=desc,
                req_type=RequirementType.LLR if "LLR" in current_id.upper() else RequirementType.HLR,
                section=current_section,
                page_or_line=f"{page_prefix} {current_line_start}",
                acceptance_criteria=self._extract_acceptance_criteria(desc),
                verification_method="TEST",
                ambiguity_status=amb_status,
                ambiguity_notes=amb_notes,
            ))

        # Fallback if no explicit headers found: parse paragraphs with 'shall' or 'must'
        if not results:
            paras = [p.strip() for p in content.split("\n\n") if p.strip()]
            for p_idx, para in enumerate(paras, start=1):
                if "shall" in para.lower() or "must" in para.lower():
                    identifier = f"REQ-{p_idx:03d}"
                    first_line = para.splitlines()[0]
                    amb_status, amb_notes = self._assess_ambiguity(para, None)
                    results.append(ExtractedRequirementDTO(
                        identifier=identifier,
                        title=first_line[:80],
                        description=para,
                        req_type=RequirementType.HLR,
                        section=current_section,
                        page_or_line=f"Paragraph {p_idx}",
                        acceptance_criteria=self._extract_acceptance_criteria(para),
                        verification_method="TEST",
                        ambiguity_status=amb_status,
                        ambiguity_notes=amb_notes,
                    ))

        return results

    def _assess_ambiguity(self, text: str, acceptance: Optional[str]) -> Tuple[str, Optional[str]]:
        combined = f"{text} {acceptance or ''}".lower()
        found = [word for word in AMBIGUOUS_KEYWORDS if re.search(r"\b" + re.escape(word) + r"\b", combined)]
        if found:
            return "AMBIGUOUS", f"Ambiguous terms detected per DO-178C guidelines: {', '.join(found)}"
        return "CLEAR", None

    def _extract_acceptance_criteria(self, text: str) -> Optional[str]:
        for line in text.splitlines():
            if "acceptance" in line.lower() or "shall return" in line.lower() or "returns" in line.lower():
                return line.strip()
        return None

    def generate_candidate_test_cases(
        self,
        identifier: str,
        description: str,
        target_function_name: Optional[str] = None,
        acceptance_criteria: Optional[str] = None,
    ) -> List[CandidateCaseDTO]:
        """Generates candidate test cases deterministically from requirement text."""
        combined_text = f"{description}\n{acceptance_criteria or ''}"

        # 1. Parse comparison expressions e.g. pressure > 900, altitude < 10000
        # Supported operators: >, <, >=, <=, ==, !=
        comp_pattern = re.compile(
            r"([a-zA-Z_][a-zA-Z0-9_]*)\s*(>|>=|<|<=|==|!=)\s*(-?\d+(?:\.\d+)?)"
        )
        conditions = comp_pattern.findall(combined_text)

        # 2. Parse return value e.g. "return 1", "shall return 0", "returns 1"
        ret_pattern = re.compile(r"(?:return|returns|output|outputs)\s*([0-9]+|true|false)", re.IGNORECASE)
        ret_match = ret_pattern.search(combined_text)
        has_explicit_return = bool(ret_match)
        expected_ret = 1
        if ret_match:
            val_str = ret_match.group(1).lower()
            expected_ret = 1 if val_str in ["1", "true"] else (0 if val_str in ["0", "false"] else int(val_str))

        # Check if function name mentioned in text
        if not target_function_name:
            fn_match = re.search(r"\b([a-zA-Z_][a-zA-Z0-9_]*_control|[a-zA-Z_][a-zA-Z0-9_]*_calc|[a-zA-Z_][a-zA-Z0-9_]*)\s*\(", combined_text)
            if fn_match:
                target_function_name = fn_match.group(1)

        cases: List[CandidateCaseDTO] = []

        if not conditions:
            # Ambiguous or non-quantitative requirement
            cases.append(CandidateCaseDTO(
                name=f"TC-{identifier}-HUMAN-REVIEW",
                case_category="NORMAL",
                rationale="Requirement contains no quantitative condition thresholds. Human authoring required.",
                target_function_name=target_function_name,
                preconditions={"requirement": identifier},
                input_vectors=[{"vector_index": 1, "inputs": {}, "expected_outputs": {}}],
                expected_outputs={},
                is_expected_result_uncertain=True,
                uncertainty_reason="Requirement lacks quantitative thresholds and explicit return value specification; requires human clarification.",
            ))
            return cases

        # We have conditions! Build nominal, boundary, invalid, and condition-based cases
        nominal_inputs: Dict[str, Any] = {}
        boundary_pass_inputs: Dict[str, Any] = {}
        boundary_fail_inputs: Dict[str, Any] = {}
        off_nominal_inputs: Dict[str, Any] = {}

        cond_specs: List[Dict[str, Any]] = []

        for var_name, op, val_str in conditions:
            val = float(val_str) if "." in val_str else int(val_str)
            cond_specs.append({"var": var_name, "op": op, "val": val})

            if op in [">", ">="]:
                step = 1 if isinstance(val, int) else 0.5
                nominal_inputs[var_name] = val + (step * 50)
                boundary_pass_inputs[var_name] = val + (step if op == ">" else 0)
                boundary_fail_inputs[var_name] = val - step
                off_nominal_inputs[var_name] = val - (step * 50)
            elif op in ["<", "<="]:
                step = 1 if isinstance(val, int) else 0.5
                nominal_inputs[var_name] = val - (step * 50)
                boundary_pass_inputs[var_name] = val - (step if op == "<" else 0)
                boundary_fail_inputs[var_name] = val + step
                off_nominal_inputs[var_name] = val + (step * 50)
            elif op == "==":
                nominal_inputs[var_name] = val
                boundary_pass_inputs[var_name] = val
                boundary_fail_inputs[var_name] = val + 1
                off_nominal_inputs[var_name] = val + 10

        # Case 1: Normal / Nominal Case
        cases.append(CandidateCaseDTO(
            name=f"TC-{identifier}-NOMINAL",
            case_category="NORMAL",
            rationale=f"Verify nominal operation when all parameters ({', '.join(nominal_inputs.keys())}) satisfy specified criteria.",
            target_function_name=target_function_name,
            preconditions={"operational_mode": "NORMAL", "requirement_id": identifier},
            input_vectors=[{"vector_index": 1, "inputs": nominal_inputs, "expected_outputs": {"return": expected_ret}}],
            expected_outputs={"return": expected_ret},
            is_expected_result_uncertain=not has_explicit_return,
            uncertainty_reason=None if has_explicit_return else "Expected return value not explicitly stated in requirement text.",
        ))

        # Case 2: Boundary Value Testing (PASS Boundary)
        cases.append(CandidateCaseDTO(
            name=f"TC-{identifier}-BOUNDARY-PASS",
            case_category="BOUNDARY",
            rationale="Verify parameter thresholds at the exact positive boundary limit.",
            target_function_name=target_function_name,
            preconditions={"operational_mode": "BOUNDARY", "requirement_id": identifier},
            input_vectors=[{"vector_index": 1, "inputs": boundary_pass_inputs, "expected_outputs": {"return": expected_ret}}],
            expected_outputs={"return": expected_ret},
            is_expected_result_uncertain=not has_explicit_return,
            uncertainty_reason=None if has_explicit_return else "Expected return value not explicitly stated in requirement text.",
        ))

        # Case 3: Boundary Value Testing (FAIL Boundary)
        fail_expected = 0 if expected_ret != 0 else 1
        cases.append(CandidateCaseDTO(
            name=f"TC-{identifier}-BOUNDARY-FAIL",
            case_category="BOUNDARY",
            rationale="Verify threshold crossing where boundary parameters exceed acceptable limits.",
            target_function_name=target_function_name,
            preconditions={"operational_mode": "BOUNDARY", "requirement_id": identifier},
            input_vectors=[{"vector_index": 1, "inputs": boundary_fail_inputs, "expected_outputs": {"return": fail_expected}}],
            expected_outputs={"return": fail_expected},
            is_expected_result_uncertain=not has_explicit_return,
            uncertainty_reason=None if has_explicit_return else "Expected return value not explicitly stated in requirement text.",
        ))

        # Case 4: Invalid / Off-Nominal Inputs
        cases.append(CandidateCaseDTO(
            name=f"TC-{identifier}-OFF-NOMINAL",
            case_category="INVALID_INPUT",
            rationale="Verify system response when operating outside specified ranges.",
            target_function_name=target_function_name,
            preconditions={"operational_mode": "OFF_NOMINAL", "requirement_id": identifier},
            input_vectors=[{"vector_index": 1, "inputs": off_nominal_inputs, "expected_outputs": {"return": fail_expected}}],
            expected_outputs={"return": fail_expected},
            is_expected_result_uncertain=not has_explicit_return,
            uncertainty_reason=None if has_explicit_return else "Expected return value not explicitly stated in requirement text.",
        ))

        # Case 5: Condition-based Independence Cases (MC/DC candidate vectors)
        # Flip each condition individually while holding others in nominal state
        if len(cond_specs) > 1:
            for idx, c in enumerate(cond_specs, start=1):
                vec_inputs = dict(nominal_inputs)
                var = c["var"]
                op = c["op"]
                val = c["val"]
                step = 1 if isinstance(val, int) else 0.5
                if op in [">", ">="]:
                    vec_inputs[var] = val - (step * 20)
                elif op in ["<", "<="]:
                    vec_inputs[var] = val + (step * 20)
                elif op == "==":
                    vec_inputs[var] = val + 5

                cases.append(CandidateCaseDTO(
                    name=f"TC-{identifier}-COND-ISOLATE-{var.upper()}",
                    case_category="CONDITION_BASED",
                    rationale=f"Isolate condition '{var} {op} {val}' to demonstrate independent effect on outcome for DO-178C MC/DC intent.",
                    target_function_name=target_function_name,
                    preconditions={"operational_mode": "CONDITION_ISOLATION", "requirement_id": identifier},
                    input_vectors=[{"vector_index": 1, "inputs": vec_inputs, "expected_outputs": {"return": fail_expected}}],
                    expected_outputs={"return": fail_expected},
                    is_expected_result_uncertain=not has_explicit_return,
                    uncertainty_reason=None if has_explicit_return else "Expected return value not explicitly stated in requirement text.",
                ))

        return cases
